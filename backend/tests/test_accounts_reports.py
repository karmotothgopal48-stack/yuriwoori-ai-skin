"""Signup/session, admin protection and the Excel report builder. No database or network needed:
protected routes are rejected before any database access, and the workbook is built from an empty report."""

import os
import sys
from datetime import timedelta
from io import BytesIO

import pytest
from fastapi import HTTPException
from pydantic import ValidationError

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core import security
from app.core.config import settings
from app.main import app
from app.schemas import EngagementEventRequest, SignupRequest
from app.services.report_excel import build_workbook

client = TestClient(app)


@pytest.fixture(autouse=True)
def _secret(monkeypatch):
    monkeypatch.setattr(settings, "jwt_secret", "x" * 48)


def test_password_hash_verifies_and_rejects():
    h = security.hash_password("correct horse battery", iterations=1000)
    assert h.startswith("pbkdf2_sha256$1000$") and "correct horse" not in h
    assert security.verify_password("correct horse battery", h)
    assert not security.verify_password("wrong password", h)
    assert not security.verify_password("x", "not-a-hash")


def test_token_roundtrip_role_and_expiry():
    claims = security.decode_token(security.create_token("abc", security.ROLE_ADMIN, timedelta(minutes=5)))
    assert claims["sub"] == "abc" and claims["role"] == "admin"
    with pytest.raises(HTTPException) as e:
        security.decode_token(security.create_token("abc", security.ROLE_USER, timedelta(seconds=-1)))
    assert e.value.status_code == 401


def test_tokens_refused_without_a_strong_secret(monkeypatch):
    monkeypatch.setattr(settings, "jwt_secret", "short")
    with pytest.raises(HTTPException) as e:
        security.create_token("abc", security.ROLE_USER, timedelta(minutes=5))
    assert e.value.status_code == 503


def test_signup_and_event_validation():
    assert SignupRequest(name="Asha", whatsapp="9876543210", email="asha@example.com", consent=True).email == "asha@example.com"
    for bad in ({"name": "", "whatsapp": "9876543210", "email": "a@b.co", "consent": True},
                {"name": "A", "whatsapp": "9876543210", "email": "not-an-email", "consent": True},
                {"name": "A", "email": "a@b.co", "consent": True}):  # WhatsApp number is required
        with pytest.raises(ValidationError):
            SignupRequest(**bad)
    with pytest.raises(ValidationError):
        EngagementEventRequest(event_type="purchase", scan_id="00000000-0000-0000-0000-000000000001")


def test_scan_start_requires_signup():
    r = client.post("/v1/scans")
    assert r.status_code == 401 and "name and email" in r.json()["detail"].lower()


def test_scan_reads_require_a_user_session():
    scan_id = "00000000-0000-0000-0000-000000000001"
    for path in (f"/v1/scans/{scan_id}", f"/v1/scans/{scan_id}/recommendations", f"/v1/scans/{scan_id}/routine",
                 f"/v1/scans/{scan_id}/compatibility", f"/v1/catalogue/recommendations?scan_id={scan_id}",
                 f"/v1/progress/{scan_id}"):
        assert client.get(path).status_code == 401, path


def test_admin_routes_need_an_admin_token():
    for path in ("/v1/admin/summary", "/v1/admin/analyses", "/v1/admin/report.xlsx"):
        assert client.get(path).status_code == 401
        user_token = security.create_token("00000000-0000-0000-0000-000000000001", security.ROLE_USER, timedelta(minutes=5))
        assert client.get(path, headers={"Authorization": f"Bearer {user_token}"}).status_code == 403


def test_admin_login_rejects_wrong_password(monkeypatch):
    monkeypatch.setattr(settings, "admin_email", "admin@example.com")
    monkeypatch.setattr(settings, "admin_password_hash", security.hash_password("right-password-123", iterations=1000))
    assert client.post("/v1/admin/login", json={"email": "admin@example.com", "password": "nope"}).status_code == 401
    ok = client.post("/v1/admin/login", json={"email": "ADMIN@example.com", "password": "right-password-123"})
    assert ok.status_code == 200 and security.decode_token(ok.json()["token"])["role"] == "admin"


def test_excel_workbook_has_required_sheets_and_columns():
    empty = {"summary": {"registered_users": 0, "scans_started": 0, "completed_analyses": 0, "unique_users_analyzed": 0,
                         "analyses_with_recommendations": 0, "products_recommended": 0, "analyses_recommendations_viewed": 0,
                         "product_clicks": 0, "analyses_with_product_click": 0, "unique_products_clicked": 0,
                         "confirmed_purchases": None, "purchase_tracking": "Not available", "funnel": [
                             {"stage": "Registered users", "count": 0, "pct_of_previous": None},
                             {"stage": "Confirmed purchases", "count": None, "pct_of_previous": None}],
                         "timezone": "IST (UTC+05:30)", "generated_at": "2026-01-01 00:00:00"},
             "analyses": [], "users": [], "recommendations": [], "clicks": [], "product_engagement": []}
    wb = load_workbook(BytesIO(build_workbook(empty)))
    assert wb.sheetnames == ["Summary", "Analyses", "Users", "Recommendations", "Product Clicks", "Notes"]
    headers = [c.value for c in wb["Analyses"][1]]
    for col in ("User Name", "Email", "Analysis ID", "Date", "Time", "IP Address", "Skin Type", "Detected Concerns",
                "Recommended Products", "Product Clicked", "Purchase Status"):
        assert col in headers
    assert wb["Analyses"].max_row == 1  # no invented rows
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows() if r[0].value}
    assert summary["Confirmed purchases"] == "Not available"


def test_health_at_root_and_v1():
    for path in ("/health", "/v1/health"):
        r = client.get(path)
        assert r.status_code == 200 and r.json()["status"] == "ok"


def test_personal_scan_data_needs_the_owner_session():
    scan_id = "00000000-0000-0000-0000-000000000001"
    # passport (skin type / concerns / score) and the budget agent's use of a scan profile
    assert client.post(f"/v1/passport/save/{scan_id}").status_code == 401
    assert client.get("/v1/passport").status_code == 401
    assert client.post("/v1/shopping-agent/build", json={"message": "2000 budget", "scan_id": scan_id}).status_code == 401


def test_coach_message_is_length_limited():
    for message in ("", "x" * 2001):
        assert client.post("/v1/coach/message", json={"message": message}).status_code == 422


def test_cors_preflight_allows_the_local_frontend_for_signup():
    # The local frontend (Vite, port 2999) calls the API from both host names; the browser preflights the
    # JSON POST. A refused preflight shows in the app as "Unable to connect to the YuriWoori server".
    for origin in ("http://localhost:2999", "http://127.0.0.1:2999"):
        r = client.options("/v1/users/signup", headers={
            "Origin": origin, "Access-Control-Request-Method": "POST", "Access-Control-Request-Headers": "content-type"})
        assert r.status_code == 200 and r.headers["access-control-allow-origin"] == origin
    r = client.options("/v1/users/signup", headers={"Origin": "https://evil.example", "Access-Control-Request-Method": "POST"})
    assert "access-control-allow-origin" not in r.headers


def test_signup_route_exists_and_validates_its_body():
    # the contract: POST /v1/users/signup {name, whatsapp, email, consent}
    assert client.post("/v1/users/signup", json={"name": "A", "whatsapp": "12", "email": "a@b.co", "consent": True}).status_code == 422
    assert client.post("/v1/users/signup", json={"name": "A", "email": "not-an-email", "consent": True}).status_code == 422
    assert client.post("/v1/users/signup", json={"email": "a@b.co", "consent": True}).status_code == 422
