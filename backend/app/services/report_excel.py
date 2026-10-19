"""Excel export of the admin report (real database rows only), using openpyxl."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.reporting import PURCHASE_STATUS

HEADER_FILL = PatternFill("solid", fgColor="005C5D")  # YuriWoori deep jade
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int] | None = None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths or [18] * len(headers), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def build_workbook(report: dict) -> bytes:
    s = report["summary"]
    wb = Workbook()
    wb.remove(wb.active)

    na = "Not available"
    summary_rows = [
        ["Total registered users", s["registered_users"]],
        ["Scans started", s["scans_started"]],
        ["Total completed analyses", s["completed_analyses"]],
        ["Unique users analyzed", s["unique_users_analyzed"]],
        ["Analyses with recommendations", s["analyses_with_recommendations"]],
        ["Product recommendations (total)", s["products_recommended"]],
        ["Analyses where recommendations were viewed", s["analyses_recommendations_viewed"]],
        ["Product clicks", s["product_clicks"]],
        ["Analyses with a product click", s["analyses_with_product_click"]],
        ["Unique products clicked", s["unique_products_clicked"]],
        ["Confirmed purchases", na],
        ["Purchase tracking", s["purchase_tracking"]],
        [],
        ["Funnel stage", "Count", "% of previous stage"],
        *[[f["stage"], na if f["count"] is None else f["count"], "" if f["pct_of_previous"] is None else f"{f['pct_of_previous']}%"]
          for f in s["funnel"]],
        [],
        ["Report generated", f"{s['generated_at']} {s['timezone']}"],
    ]
    ws = _sheet(wb, "Summary", ["Metric", "Value"], summary_rows, [44, 30, 20])
    funnel_header = summary_rows.index(["Funnel stage", "Count", "% of previous stage"])
    for cell in ws[funnel_header + 2]:  # +1 for the header row, +1 because sheet rows start at 1
        cell.font = Font(bold=True)

    _sheet(wb, "Analyses",
           ["User Name", "Email", "Analysis ID", "Date", "Time", "IP Address", "Skin Type", "Overall Score", "Detected Concerns",
            "Recommended Products", "Recommendations Viewed", "Product Clicked", "Clicked Products", "Purchase Status"],
           [[a["user_name"], a["email"], a["analysis_id"], a["date"], a["time"], a["ip_address"], a["skin_type"], a["overall_score"],
             ", ".join(a["detected_concerns"]) or "None detected", "\n".join(a["recommended_products"]) or "None",
             "Yes" if a["recommendations_shown"] else "No", "Yes" if a["product_clicked"] else "No",
             "\n".join(a["clicked_products"]), a["purchase_status"]] for a in report["analyses"]],
           [18, 26, 38, 12, 10, 16, 12, 10, 24, 44, 14, 12, 40, 30])
    _sheet(wb, "Users", ["Name", "Email", "Registered At", "Consent Version", "Completed Analyses"],
           [[u["name"], u["email"], u["registered_at"], u["consent_version"], u["analyses"]] for u in report["users"]],
           [22, 30, 20, 16, 18])
    _sheet(wb, "Recommendations", ["Analysis ID", "Email", "Rank", "Product"],
           [[r["analysis_id"], r["email"], r["rank"], r["product"]] for r in report["recommendations"]], [38, 30, 8, 52])
    _sheet(wb, "Product Clicks", ["Analysis ID", "Email", "Product", "Product URL", "Clicked At", "IP Address"],
           [[c["analysis_id"], c["email"], c["product"], c["product_url"], c["clicked_at"], c["ip_address"]] for c in report["clicks"]],
           [38, 30, 44, 56, 20, 16])
    _sheet(wb, "Notes", ["Note"], [
        [f"All times are {s['timezone']}."],
        ["Only signed-up users (name, email, accepted privacy notice) are included."],
        ["IP address is the server-side request address when the scan was started (behind a proxy it requires trusted proxy headers)."],
        ["Recommended Products = the ranked list served to the user for that analysis, from the official YuriWoori catalogue."],
        ["Recommendations Viewed = the user opened the Products screen for that analysis."],
        ["Product Clicked = the user opened an official yuriwoori.com product page from a recommendation or routine step."],
        [f"Purchase Status = '{PURCHASE_STATUS}'. A product click is NOT a purchase. Confirming purchases needs a Shopify "
         "order integration (Admin API read_orders access or an orders/create webhook) that is not configured."],
    ], [120])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
